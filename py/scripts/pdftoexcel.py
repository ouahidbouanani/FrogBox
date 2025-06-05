import tempfile
import io
from pdf2docx import Converter
from docx import Document
import openpyxl

def convert_pdf_to_excel_function(file):
    # Lire le PDF envoyé dans la requête
    pdf_file = io.BytesIO(file.read())

    # Créer un fichier temporaire pour sauvegarder le contenu PDF
    with tempfile.NamedTemporaryFile(delete=False, mode='wb') as temp_pdf:
        temp_pdf.write(pdf_file.read())
        temp_pdf_path = temp_pdf.name  # Obtenir le chemin du fichier temporaire

    # Convertir le PDF en un fichier DOCX
    docx_temp_path = temp_pdf_path.replace('.pdf', '.docx')
    cv = Converter(temp_pdf_path)
    cv.convert(docx_temp_path, start=0, end=None)

    # Ouvrir le fichier DOCX converti
    doc = Document(docx_temp_path)

    # Créer un fichier Excel temporaire
    wb = openpyxl.Workbook()
    ws = wb.active
    row_index = 1

    # Parcourir le fichier DOCX pour extraire les données et les convertir en Excel
    for para in doc.paragraphs:
        if para.text.strip():
            ws.append([para.text])  # Ajouter une nouvelle ligne pour chaque paragraphe de texte

    # Sauvegarder l'Excel dans un fichier temporaire
    with tempfile.NamedTemporaryFile(delete=False, mode='wb') as temp_excel:
        output_excel_path = temp_excel.name
        wb.save(output_excel_path)

    return output_excel_path
